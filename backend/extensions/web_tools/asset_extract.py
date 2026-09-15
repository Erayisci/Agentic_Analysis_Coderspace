"""Bounded PDF, spreadsheet, DOCX and image extraction inside a killable child."""

import csv
from contextlib import contextmanager
from datetime import date, datetime
import io
import math
from pathlib import Path
import subprocess
import warnings
import zipfile

from .asset_common import AssetFailure, WARNING, timestamp


class Collector:
    def __init__(self, request, metadata):
        self.maximum = request["max_chars"]
        self.parts, self.sections, self.warnings = [], [], [WARNING]
        self.original, self.returned, self.limited = 0, 0, False
        self.result = {"requested_url": request["url"], **metadata, "fetched_at": timestamp(),
                       "kind": request["kind"], "source_trust": "untrusted_external"}

    def add(self, text, location, method, rows=None):
        text = text.strip()
        if not text:
            return
        heading = ("\n\n" if self.parts else "") + "[" + location + "]\n"
        rendered = heading + text
        self.original += len(rendered)
        room = self.maximum - self.returned
        if len(rendered) > room:
            self.limited = True
        if room <= len(heading):
            return
        selected = text[:room - len(heading)]
        self.parts.append(heading + selected)
        self.returned += len(heading + selected)
        section = {"location": location, "method": method, "text": selected}
        if rows is not None and len(rendered) <= room:
            section["rows"] = rows
        self.sections.append(section)

    def limit(self, message):
        self.limited = True
        if message not in self.warnings:
            self.warnings.append(message)

    def finish(self):
        content = "".join(self.parts)
        return {**self.result, "status": "partial" if self.limited else ("ok" if content else "empty"),
                "content": content, "sections": self.sections, "original_chars": self.original,
                "returned_chars": len(content), "max_chars": self.maximum, "truncated": self.limited,
                "warnings": self.warnings, "error": None, "cache": {"hit": False, "age_seconds": 0}}


def inspect_zip(path, config):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 2000 or sum(item.file_size for item in members) > config.asset_max_archive_bytes:
            raise AssetFailure("asset_too_large")
        for item in members:
            if item.flag_bits & 1 or item.file_size > max(1, item.compress_size) * 200:
                raise AssetFailure("asset_too_large")
        return {item.filename for item in members}


def detect(path, metadata, config):
    with open(path, "rb") as stream:
        signature = stream.read(1024)
    if signature.startswith(b"%PDF-"):
        return "pdf"
    if signature.startswith(b"PK"):
        names = inspect_zip(path, config)
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
    if signature.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        return "xls"
    if signature.startswith((b"\x89PNG\r\n", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*")) or (
            signature.startswith(b"RIFF") and signature[8:12] == b"WEBP"):
        return "image"
    media = metadata.get("content_type", "")
    if media in {"text/csv", "application/csv", "text/tab-separated-values"} or (
            media in {"text/plain", "application/octet-stream"} and metadata["final_url"].split("?", 1)[0].lower().endswith(".csv")):
        if b"<html" in signature.lower() or b"<!doctype" in signature.lower() or b"\x00" in signature[:10] and not signature.startswith((b"\xff\xfe", b"\xfe\xff")):
            raise AssetFailure("unsupported_content_type")
        return "csv"
    raise AssetFailure("unsupported_content_type")


def cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)[:512]


def add_rows(collector, rows, location, config, method="table"):
    selected = []
    size = 0
    for index, row in enumerate(rows):
        if index >= config.asset_max_rows:
            collector.limit("The configured row limit was reached.")
            break
        values = list(row)
        if any(len(str(value)) > 512 for value in values[:config.asset_max_columns] if value is not None):
            collector.limit("Some cell values exceeded the 512-character cell limit.")
        if len(values) > config.asset_max_columns:
            collector.limit("The configured column limit was reached.")
        values = [cell(value) for value in values[:config.asset_max_columns]]
        size += sum(len(value) for value in values)
        selected.append(values)
        if size >= collector.maximum:
            collector.limit("The configured text limit was reached.")
            break
    collector.add("\n".join("\t".join(row) for row in selected), location, method, rows=selected)


def png_bytes(image, config):
    from PIL import Image
    image = image.convert("RGB")
    image.thumbnail((config.asset_image_edge, config.asset_image_edge), Image.Resampling.LANCZOS)
    output = io.BytesIO()
    image.save(output, format="PNG")
    if output.tell() > 4 * 1024 * 1024:
        raise AssetFailure("asset_too_large")
    return output.getvalue()


class Interpreter:
    def __init__(self, request, config, proxy, collector):
        self.request, self.config, self.proxy, self.collector = request, config, proxy, collector
        self.calls, self.ocr_pages = 0, 0

    def model(self, images, locations, *, ocr):
        from backend.model_clients.kloudeks import KloudeksClient, ModelFailure
        from .asset_cache import AssetStore
        if self.calls >= self.config.model_max_calls_per_read:
            self.collector.limit("The per-read model call limit was reached.")
            return False
        if not self.config.kloudeks_api_key:
            raise AssetFailure("model_not_configured")
        AssetStore(self.config).consume_model_call()
        self.calls += 1
        client = KloudeksClient(self.config.kloudeks_base_url, self.config.kloudeks_api_key, self.proxy,
                               timeout=min(60, self.config.asset_timeout_seconds))
        try:
            result = client.interpret(images, model=self.config.kloudeks_ocr_model if ocr else self.config.kloudeks_vision_model,
                                      max_tokens=self.config.model_max_tokens, question=self.request["question"], ocr=ocr)
        except ModelFailure as error:
            raise AssetFailure(error.code) from None
        self.collector.add(result["text"], ", ".join(locations), "mia_ocr" if ocr else "mia_vision")
        self.collector.result["model"] = result["model"]
        self.collector.warnings.append("Model-derived text may contain recognition or interpretation errors; verify against the source.")
        if result["truncated"]:
            self.collector.limit("The model output token limit was reached.")
        return True

    def ocr(self, image, location):
        if self.ocr_pages >= self.config.ocr_max_pages:
            self.collector.limit("The configured OCR page limit was reached.")
            return
        data = png_bytes(image, self.config)
        if self.config.ocr_provider == "kloudeks":
            if self.model([data], [location], ocr=True):
                self.ocr_pages += 1
            return
        result = subprocess.run(["tesseract", "stdin", "stdout", "-l", self.config.ocr_languages, "--psm", "3"],
                                input=data, capture_output=True, timeout=min(30, self.config.asset_timeout_seconds))
        if result.returncode:
            raise AssetFailure("parse_error")
        self.ocr_pages += 1
        text = result.stdout.decode("utf-8", errors="replace")
        self.collector.add(text, location, "local_ocr")
        self.collector.warnings.append("OCR text can misread numbers and layout; verify against the source.")


def render_pdf(document, index, config):
    page = document[index]
    try:
        width, height = page.get_size()
        if width <= 0 or height <= 0:
            raise AssetFailure("parse_error")
        # PDFium rounds each rendered dimension up. Leave one pixel of room in
        # each dimension so the configured pixel budget also bounds PDF scans.
        scale = min(2.0, (config.asset_image_edge - 1) / max(width, height),
                    (math.sqrt(config.asset_max_image_pixels) - 1) / max(width, height))
        bitmap = page.render(scale=scale)
        try:
            return bitmap.to_pil().copy()
        finally:
            bitmap.close()
    finally:
        page.close()


def extract_pdf(path, request, config, collector, interpreter):
    from pypdf import PdfReader
    import pdfplumber
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise AssetFailure("parse_error")
    total = len(reader.pages)
    start = request["start_page"] - 1
    if start >= total:
        raise AssetFailure("invalid_request")
    end = min(total, start + request["max_pages"])
    collector.result.update({"total_pages": total, "pages_processed": end - start})
    if start > 0 or end < total:
        collector.limit("Only the requested, bounded page range was processed.")
    rendered = None
    vision_images, vision_locations, ocr_images, ocr_locations = [], [], [], []
    try:
        with pdfplumber.open(path) as tables:
            for index in range(start, end):
                location = f"Page {index + 1}"
                text = reader.pages[index].extract_text(extraction_mode="layout") or ""
                collector.add(text, location, "pdf_text")
                if collector.returned < collector.maximum:
                    table = tables.pages[index].extract_table()
                    if table:
                        add_rows(collector, table, location + ", table", config)
                needs_ocr = request["ocr"] and not text.strip()
                if needs_ocr and interpreter.ocr_pages + len(ocr_images) >= config.ocr_max_pages:
                    collector.limit("The configured OCR page limit was reached.")
                    needs_ocr = False
                needs_vision = request["vision"] and len(vision_images) < config.asset_max_images
                if needs_ocr or needs_vision:
                    if rendered is None:
                        import pypdfium2
                        rendered = pypdfium2.PdfDocument(path)
                    with render_pdf(rendered, index, config) as image:
                        if needs_ocr:
                            if config.ocr_provider == "kloudeks":
                                ocr_images.append(png_bytes(image, config))
                                ocr_locations.append(location)
                            else:
                                interpreter.ocr(image, location)
                        if needs_vision:
                            vision_images.append(png_bytes(image, config))
                            vision_locations.append(location)
                elif not text.strip() and not request["ocr"]:
                    collector.limit(f"{location} has no extracted text; enable OCR for scanned content.")
                tables.pages[index].close()
        for offset in range(0, len(ocr_images), 3):
            batch = ocr_images[offset:offset + 3]
            if not interpreter.model(batch, ocr_locations[offset:offset + 3], ocr=True):
                break
            interpreter.ocr_pages += len(batch)
        if vision_images:
            interpreter.model(vision_images, vision_locations, ocr=False)
            if end - start > len(vision_images):
                collector.limit("The configured vision image limit was reached.")
    finally:
        if rendered is not None:
            rendered.close()


@contextmanager
def bounded_image(path, config):
    from PIL import Image
    previous_limit = Image.MAX_IMAGE_PIXELS
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        Image.MAX_IMAGE_PIXELS = config.asset_max_image_pixels
        try:
            with Image.open(path) as source:
                yield source
        except (Image.DecompressionBombError, Image.DecompressionBombWarning):
            raise AssetFailure("asset_too_large") from None
        finally:
            Image.MAX_IMAGE_PIXELS = previous_limit


def extract_image(path, request, config, collector, interpreter):
    from PIL import ImageOps
    with bounded_image(path, config) as source:
        if source.width * source.height > config.asset_max_image_pixels:
            raise AssetFailure("asset_too_large")
        collector.result["image"] = {"width": source.width, "height": source.height, "format": source.format}
        source.seek(0)
        with ImageOps.exif_transpose(source) as image:
            if request["ocr"]:
                interpreter.ocr(image, "Image 1")
            if request["vision"]:
                interpreter.model([png_bytes(image, config)], ["Image 1"], ocr=False)
        if getattr(source, "n_frames", 1) > 1:
            collector.limit("Only the first image frame was processed.")
    if not request["ocr"] and not request["vision"]:
        collector.warnings.append("Image metadata only: enable OCR or request vision to extract meaning.")


def extract(path, request, metadata, config, proxy):
    kind = detect(path, metadata, config)
    if (request["kind"] == "image") != (kind == "image"):
        raise AssetFailure("unsupported_content_type")
    collector = Collector(request, {**metadata, "format": kind})
    interpreter = Interpreter(request, config, proxy, collector)
    if kind == "pdf":
        extract_pdf(path, request, config, collector, interpreter)
    elif kind == "image":
        extract_image(path, request, config, collector, interpreter)
    elif kind == "csv":
        raw = Path(path).read_bytes()
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = raw.decode(encoding)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(1048576)
        add_rows(collector, csv.reader(io.StringIO(text), dialect), "CSV rows starting at 1", config)
    elif kind == "xlsx":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(Path(path).read_bytes()), read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.sheetnames) > config.asset_max_sheets:
                collector.limit("The configured sheet limit was reached.")
            for sheet in book.worksheets[:config.asset_max_sheets]:
                if (sheet.max_row or 0) > config.asset_max_rows or (sheet.max_column or 0) > config.asset_max_columns:
                    collector.limit("The configured row/column limit was reached.")
                add_rows(collector, sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or config.asset_max_rows,
                          config.asset_max_rows), max_col=min(sheet.max_column or config.asset_max_columns,
                          config.asset_max_columns), values_only=True), "Sheet " + sheet.title + ", from A1", config)
            collector.warnings.append("Spreadsheet formulas are not executed; cached formula values may be missing or stale.")
        finally:
            book.close()
    elif kind == "xls":
        import xlrd
        with xlrd.open_workbook(path, on_demand=True) as book:
            if book.nsheets > config.asset_max_sheets:
                collector.limit("The configured sheet limit was reached.")
            for index in range(min(book.nsheets, config.asset_max_sheets)):
                sheet = book.sheet_by_index(index)
                def row_values(row):
                    for item in sheet.row(row, 0, min(sheet.ncols, config.asset_max_columns)):
                        if item.ctype == xlrd.XL_CELL_DATE:
                            yield xlrd.xldate_as_datetime(item.value, book.datemode)
                        elif item.ctype == xlrd.XL_CELL_BOOLEAN:
                            yield bool(item.value)
                        elif item.ctype == xlrd.XL_CELL_ERROR:
                            yield xlrd.error_text_from_code.get(item.value, "#ERROR")
                        else:
                            yield item.value
                rows = (row_values(row) for row in range(min(sheet.nrows, config.asset_max_rows + 1)))
                add_rows(collector, rows, "Sheet " + sheet.name + ", from A1", config)
                if sheet.ncols > config.asset_max_columns:
                    collector.limit("The configured column limit was reached.")
            collector.warnings.append("Spreadsheet formulas are not executed; cached formula values may be missing or stale.")
    elif kind == "docx":
        from defusedxml import ElementTree
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        body = root.find("w:body", namespace)
        for index, element in enumerate(body if body is not None else []):
            if collector.returned >= collector.maximum:
                collector.limit("The configured text limit was reached.")
                break
            if element.tag.endswith("}tbl"):
                rows = [["".join(cell_node.itertext()) for cell_node in row.findall("w:tc", namespace)]
                        for row in element.findall("w:tr", namespace)]
                add_rows(collector, rows, f"DOCX block {index + 1}", config)
            else:
                text = "".join(node.text or "" for node in element.findall(".//w:t", namespace))
                collector.add(text, f"DOCX block {index + 1}", "docx_text")
    collector.result["model_calls"] = interpreter.calls
    collector.result["ocr_pages"] = interpreter.ocr_pages
    return collector.finish()
